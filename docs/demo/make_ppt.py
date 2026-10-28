# -*- coding: utf-8 -*-
"""CampusHub 软件体系架构大作业 答辩PPT 生成脚本"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

SCREEN = r"C:\Users\19465\Doubao\chats\2026-09-21\new-chat\CampusHub\docs\demo\screenshots"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BLUE = RGBColor(0x30, 0x41, 0x56)
ACCENT = RGBColor(0x40, 0x9E, 0xFF)
DARK = RGBColor(0x30, 0x30, 0x30)

def add_title_slide(title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid(); bg.fill.fore_color.rgb = BLUE
    bg.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.3), Inches(1.6))
    tf = tb.text_frame; tf.text = title
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].font.size = Pt(40); tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    sb = slide.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11.3), Inches(1.2))
    stf = sb.text_frame; stf.text = subtitle
    stf.paragraphs[0].alignment = PP_ALIGN.CENTER
    stf.paragraphs[0].font.size = Pt(20)
    stf.paragraphs[0].font.color.rgb = RGBColor(0xBF, 0xCB, 0xD9)

def add_content_slide(title, bullets, image=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # title bar
    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.9))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.12), Inches(12), Inches(0.7))
    tf = tb.text_frame; tf.text = title
    tf.paragraphs[0].font.size = Pt(26); tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    # bullets
    body = slide.shapes.add_textbox(Inches(0.7), Inches(1.2), Inches(6.3), Inches(5.8))
    tf2 = body.text_frame; tf2.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        p.text = ("• " if not b.startswith("  ") else "") + b
        p.font.size = Pt(16); p.font.color.rgb = DARK
        p.space_after = Pt(8)
    if image:
        try:
            pic = slide.shapes.add_picture(image, Inches(7.2), Inches(1.3), width=Inches(5.6))
        except Exception:
            pass

# ============ 幻灯片 ============
add_title_slide("CampusHub 校园服务中心",
                "软件体系架构课程大作业 · 五视图架构设计与实现\n研一 · 长春理工大学")

add_content_slide("1. 项目背景与目标", [
    "校园信息分散：社团靠海报、竞赛靠群聊、二手靠朋友圈、失物靠公告栏 → 信息孤岛",
    "目标：构建一站式校园信息共享中心，聚合社团/竞赛/活动/二手/失物/消息 6 大类信息",
    "服务对象：学生、社团负责人、管理员三类角色",
    "规模：9 个业务模块、19 张表、5 个核心状态机、40+ REST 接口",
])

add_content_slide("2. 需求概述：角色与功能", [
    "角色（RBAC）：ADMIN 管理员 / CLUB_LEADER 社团负责人 / STUDENT 学生",
    "M1 用户权限 · M2 社团管理（审批流）· M3 竞赛组队 · M4 活动组织（报名/签到）",
    "M5 消息通知（异步）· M6 二手交易 · M7 失物招领 · M8 文件管理 · M9 数据看板",
    "核心状态机：社团 PENDING→APPROVED/REJECTED；活动 DRAFT→…→FINISHED/CANCELLED；二手 ON_SALE→RESERVED→SOLD；失物 PENDING→CLAIMED/RESOLVED",
])

add_content_slide("3. 技术栈", [
    "后端：Java 17 + Spring Boot 3.2 + MyBatis-Plus + Spring Security/JWT",
    "存储：MySQL 8（19 表）+ Redis 7（并发计数）",
    "消息：RabbitMQ 3.12（异步通知 + 死信队列）",
    "前端：Vue 3 + Element Plus + Vite + Pinia",
    "文档：SpringDoc OpenAPI（Swagger UI）· 部署：Docker Compose（中间件编排）",
])

add_content_slide("4. 五视图架构总览", [
    "逻辑视图 L1-L7：上下文 / 功能 / 容器 / 分层 / 领域模型 / 核心接口 / 关键时序",
    "开发视图 V1-V8：代码结构 / 模块依赖 / 公共层 / 开发规范",
    "运行视图 R1-R7：JWT 认证 / RBAC 授权 / Redis 并发 / MQ 异步 / 定时任务 / 文件 / 看板",
    "物理视图 P1/P4：单机部署拓扑 + Docker 编排",
    "场景视图：报名活动 / 创建社团 / 二手交易 / 失物认领 / 竞赛组队",
])

add_content_slide("5. 逻辑视图：分层架构", [
    "表现层：Vue SPA（views + router + api）",
    "接口层：Controller（/api/**），JwtAuthInterceptor 鉴权 + @RequireRole 角色切面",
    "业务层：Service（状态机 + 规则校验），EventPublisher → RabbitMQ，Redis 计数",
    "数据层：MyBatis-Plus Mapper（BaseMapper + 分页插件）",
    "基础设施：MySQL / Redis / RabbitMQ / 本地文件系统",
])

add_content_slide("6. 关键设计：活动报名并发控制", [
    "场景：多人同时报名，名额不超卖、一人一报",
    "方案：Redis INCR 原子扣减（O(1) 无锁）→ 超限 DECR 回滚拒绝",
    "兜底：DB 唯一键 uk_activity_user 保证一人一报",
    "报名成功发布 REGISTRATION_EVENT → MQ 异步写站内通知，主流程不阻塞",
], image=SCREEN + r"\12-活动详情.png")

add_content_slide("7. 运行视图：异步消息链路", [
    "直连交换机 campushub.exchange + 3 个业务队列（registration/club/trade）+ 死信队列 DLX",
    "DomainEvent 统一事件结构，消费者按 routingKey 分流",
    "内存幂等（已处理 key 集合）防重复消费",
    "通知失败不影响主流程（解耦），数据由 DB 唯一键最终一致兜底",
])

add_content_slide("8. 运行视图：定时任务", [
    "活动状态自动推进：每分钟按时间 DRAFT→PUBLISHED→REGISTERING→IN_PROGRESS→FINISHED",
    "每日 00:05 统计：stat_daily 每日一行（uk_date 幂等）",
    "看板聚合：管理员查看近 N 天新增趋势",
], image=SCREEN + r"\08-数据看板.png")

add_content_slide("9. 物理视图：部署拓扑", [
    "演示环境：本地 Windows 单机进程部署",
    "Spring Boot 应用（8080）+ Vue 开发服务器（5173，代理 /api）",
    "中间件：MySQL（3306）/ Redis（6379）/ RabbitMQ（5672）",
    "docker-compose.yml 一键编排中间件（含健康检查 + 数据卷）",
])

add_content_slide("10. 架构决策记录 ADR（8 条）", [
    "ADR-001 后端 Spring Boot 3 + Java 17",
    "ADR-002 前后端分离 Vue 3 + Vite",
    "ADR-003 MySQL 8 + MyBatis-Plus",
    "ADR-004 Redis 缓存与并发计数",
    "ADR-005 JWT + 拦截器 + RBAC 切面认证授权",
    "ADR-006 RabbitMQ 直连交换机 + 死信队列异步通知",
    "ADR-007 单机进程 + Docker 编排中间件",
    "ADR-008 按业务领域 9 模块 + 公共层划分",
], image=SCREEN + r"\11-Swagger接口文档.png")

add_content_slide("11. 一致性追溯与质量属性", [
    "需求-架构追溯矩阵 15 条：需求 → 视图 → ADR → 实现位置 全链路可追溯",
    "三个一致性问题：模块解耦边界（事件总线）/ Redis 与 DB 一致性（计数+唯一键）/ 单体演进微服务预留",
    "质量属性：性能（Redis 原子计数）、安全（JWT+MD5）、可用性（MQ 解耦）、可维护性（模块化+Skill 生成）",
])

add_content_slide("12. 开发与测试", [
    "代码规模：9 模块 40+ 接口，Java 20+ 类，Vue 11 个页面",
    "单元测试 17 个用例（JWT / 社团 / 活动 / 失物），覆盖状态机与并发回滚",
    "前端 npm run build 通过；后端 mvn compile + mvn test 通过",
    "测试账号：admin/123456（ADMIN）、student/123456（STUDENT）",
    "Git 提交 30+ 次，全程版本管理",
], image=SCREEN + r"\02-社团列表.png")

add_content_slide("13. 演示：核心页面", [
    "首页 / 社团管理 / 竞赛组队 / 活动管理 / 二手闲置 / 失物招领",
    "消息中心（异步通知）/ 数据看板 / 个人信息 / 注册登录",
    "Swagger 接口文档自动生成",
], image=SCREEN + r"\01-首页.png")

add_content_slide("14. 总结与展望", [
    "完成：五视图设计 + 8 ADR + 全栈实现 + 单测 + 文档 + Skill 包",
    "架构亮点：状态机驱动业务、Redis 并发控制、MQ 异步解耦、JWT+RBAC 安全",
    "展望：消息中心拆分微服务、部署容器化、引入 CI/CD 与监控",
])

prs.save(r"C:\Users\19465\Doubao\chats\2026-09-21\new-chat\CampusHub\docs\demo\CampusHub-答辩PPT.pptx")
print("PPT saved")
